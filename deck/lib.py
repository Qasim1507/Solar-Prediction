"""Shared layout helpers for the defence deck."""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_AUTO_SIZE

NAVY  = RGBColor(0x1F,0x3B,0x63)
ACC   = RGBColor(0xC4,0x62,0x2D)
TEAL  = RGBColor(0x2E,0x7D,0x74)
RED   = RGBColor(0xB3,0x28,0x2D)
INK   = RGBColor(0x25,0x2F,0x3D)
GREY  = RGBColor(0x6B,0x72,0x80)
LGREY = RGBColor(0xD8,0xDE,0xE8)
BG2   = RGBColor(0xF3,0xF5,0xF9)
WHITE = RGBColor(0xFF,0xFF,0xFF)
BODY  = "Calibri"
FOOT  = "Satellite–Weather Fusion for Short-Horizon GHI Forecasting  ·  M.Eng. Computer Engineering"

W, H = 13.333, 7.5

def new_deck():
    prs = Presentation()
    prs.slide_width  = Inches(W)
    prs.slide_height = Inches(H)
    return prs

def _tf(shape, size=14, color=INK, bold=False, space_after=6, align=PP_ALIGN.LEFT):
    tf = shape.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    return tf

def txt(slide, x, y, w, h, text, size=14, color=INK, bold=False, italic=False,
        align=PP_ALIGN.LEFT, font=BODY, anchor=MSO_ANCHOR.TOP, spacing=1.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = spacing
    r = p.add_run(); r.text = text
    r.font.size = Pt(size); r.font.bold = bold; r.font.italic = italic
    r.font.color.rgb = color; r.font.name = font
    return tb

def rect(slide, x, y, w, h, fill=None, line=None, lw=0.75, shape=MSO_SHAPE.RECTANGLE,
         adj=None):
    s = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    else:
        s.fill.solid(); s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line; s.line.width = Pt(lw)
    s.shadow.inherit = False
    s.text_frame.word_wrap = True
    if adj is not None:
        try: s.adjustments[0] = adj
        except Exception: pass
    return s

def line(slide, x1, y1, x2, y2, color=LGREY, lw=1.0, dash=None):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw)
    return c

def arrow(slide, x1, y1, x2, y2, color=GREY, lw=1.25):
    from pptx.enum.shapes import MSO_CONNECTOR
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1),
                                   Inches(x2), Inches(y2))
    c.line.color.rgb = color; c.line.width = Pt(lw)
    el = c.line._get_or_add_ln()
    from pptx.oxml.ns import qn
    import copy
    tail = el.makeelement(qn('a:tailEnd'), {'type':'triangle','w':'sm','len':'sm'})
    el.append(tail)
    return c

def slide_base(prs, title, kicker=None, number=None, rule=ACC):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = WHITE
    y = 0.34
    if kicker:
        txt(s, 0.55, y, 12.2, 0.25, kicker.upper(), size=10.5, color=ACC, bold=True)
        y += 0.28
    tsize = 24 if len(title) <= 62 else (21.5 if len(title) <= 78 else 19.5)
    txt(s, 0.55, y, 12.2, 0.5, title, size=tsize, color=NAVY, bold=True)
    line(s, 0.55, y+0.53, 12.78, y+0.53, color=LGREY, lw=1.25)
    line(s, 0.55, y+0.53, 1.55, y+0.53, color=rule, lw=2.5)
    txt(s, 0.55, 7.02, 9.0, 0.25, FOOT, size=8.5, color=GREY)
    if number is not None:
        txt(s, 11.6, 7.02, 1.18, 0.25, str(number), size=9.5, color=GREY,
            align=PP_ALIGN.RIGHT)
    return s

def bullets(slide, x, y, w, h, items, size=13.5, spacing=1.06, gap=7):
    """items: list of (level:int, text:str) or (level, text, dict(bold=,color=,size=))."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    first = True
    for it in items:
        lvl, text = it[0], it[1]
        opt = it[2] if len(it) > 2 else {}
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = spacing
        p.space_after = Pt(opt.get('after', gap))
        p.space_before = Pt(opt.get('before', 0))
        indent = {0:0.0, 1:0.22, 2:0.44}[lvl]
        p.left_indent = Inches(indent)
        marks = {0:'', 1:'▪  ', 2:'–  '}
        pre = opt.get('mark', marks[lvl])
        segs = text if isinstance(text, list) else [(text, {})]
        if pre:
            r = p.add_run(); r.text = pre
            r.font.size = Pt(opt.get('size', size)); r.font.name = BODY
            r.font.color.rgb = opt.get('markcolor', ACC if lvl == 1 else GREY)
        for seg_text, seg_opt in segs:
            r = p.add_run(); r.text = seg_text
            r.font.size = Pt(seg_opt.get('size', opt.get('size', size)))
            r.font.bold = seg_opt.get('bold', opt.get('bold', False))
            r.font.italic = seg_opt.get('italic', opt.get('italic', False))
            r.font.color.rgb = seg_opt.get('color', opt.get('color', INK))
            r.font.name = seg_opt.get('font', BODY)
    return tb

def table(slide, x, y, w, headers, rows, col_w=None, size=10.5, hsize=10.5,
          row_h=0.26, head_h=0.3, highlight=None, hl_color=None, align=None,
          zebra=True):
    nr, nc = len(rows) + 1, len(headers)
    gt = slide.shapes.add_table(nr, nc, Inches(x), Inches(y), Inches(w),
                                Inches(head_h + row_h * len(rows))).table
    gt.first_row = False
    try:
        gt._tbl.tblPr.remove(gt._tbl.tblPr.find(
            '{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId'))
    except Exception:
        pass
    if col_w:
        tot = sum(col_w)
        for i, cw in enumerate(col_w):
            gt.columns[i].width = Emu(int(Inches(w) * cw / tot))
    gt.rows[0].height = Inches(head_h)
    for i in range(len(rows)):
        gt.rows[i+1].height = Inches(row_h)
    hl = highlight or []
    for c, htext in enumerate(headers):
        cell = gt.cell(0, c)
        cell.fill.solid(); cell.fill.fore_color.rgb = NAVY
        cell.margin_left = cell.margin_right = Inches(0.06)
        cell.margin_top = cell.margin_bottom = Inches(0.02)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = cell.text_frame.paragraphs[0]
        p.alignment = (align[c] if align else PP_ALIGN.LEFT)
        r = p.add_run(); r.text = htext
        r.font.size = Pt(hsize); r.font.bold = True; r.font.color.rgb = WHITE
        r.font.name = BODY
    for i, row in enumerate(rows):
        for c, val in enumerate(row):
            cell = gt.cell(i+1, c)
            cell.margin_left = cell.margin_right = Inches(0.06)
            cell.margin_top = cell.margin_bottom = Inches(0.01)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            if i in hl:
                cell.fill.fore_color.rgb = hl_color or RGBColor(0xFD,0xF0,0xE7)
            else:
                cell.fill.fore_color.rgb = BG2 if (zebra and i % 2 == 0) else WHITE
            p = cell.text_frame.paragraphs[0]
            p.alignment = (align[c] if align else PP_ALIGN.LEFT)
            txt_, opt = (val, {}) if not isinstance(val, tuple) else val
            r = p.add_run(); r.text = str(txt_)
            r.font.size = Pt(opt.get('size', size))
            r.font.bold = opt.get('bold', i in hl)
            r.font.color.rgb = opt.get('color', INK)
            r.font.name = BODY
    return gt

def caption(slide, x, y, w, text, size=9.5, color=GREY, italic=True):
    return txt(slide, x, y, w, 0.24, text, size=size, color=color, italic=italic)

def note(slide, text):
    slide.notes_slide.notes_text_frame.text = text.strip()

def panel(slide, x, y, w, h, fill=BG2, line_c=LGREY):
    return rect(slide, x, y, w, h, fill=fill, line=line_c, lw=0.75)

def tag(slide, x, y, w, h, text, color=NAVY, fill=None, size=10):
    r = rect(slide, x, y, w, h, fill=fill or RGBColor(0xEC,0xEF,0xF4), line=None)
    tf = r.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.06)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    run = p.add_run(); run.text = text
    run.font.size = Pt(size); run.font.color.rgb = color; run.font.name = BODY
    run.font.bold = True
    return r

def box(slide, x, y, w, h, title_, lines=None, fill=WHITE, edge=NAVY, tcolor=NAVY,
        tsize=11, lsize=9.5, lcolor=GREY, align=PP_ALIGN.CENTER):
    s = rect(slide, x, y, w, h, fill=fill, line=edge, lw=1.0)
    tf = s.text_frame; tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(0.05)
    tf.margin_top = Inches(0.04); tf.margin_bottom = Inches(0.03)
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    p = tf.paragraphs[0]; p.alignment = align; p.line_spacing = 0.95
    r = p.add_run(); r.text = title_
    r.font.size = Pt(tsize); r.font.bold = True; r.font.color.rgb = tcolor
    r.font.name = BODY
    for ln in (lines or []):
        p = tf.add_paragraph(); p.alignment = align; p.line_spacing = 0.95
        p.space_before = Pt(1)
        r = p.add_run(); r.text = ln
        r.font.size = Pt(lsize); r.font.color.rgb = lcolor; r.font.name = BODY
    return s

def pic(slide, x, y, path, w=None, h=None):
    kw = {}
    if w: kw['width'] = Inches(w)
    if h: kw['height'] = Inches(h)
    return slide.shapes.add_picture(path, Inches(x), Inches(y), **kw)
