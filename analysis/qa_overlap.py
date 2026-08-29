from pptx import Presentation
EM=914400
P=Presentation('deck/GHI_Forecasting_Defence.pptx')
def est_h(sh, w):
    tot=0
    for p in sh.text_frame.paragraphs:
        t=''.join(r.text for r in p.runs)
        if not t: continue
        sz=max([(r.font.size.pt if r.font.size else 12) for r in p.runs] or [12])
        cpl=max(6,int(w*72/(sz*0.50)))
        n=max(1,-(-len(t)//cpl))
        tot+=n*sz*1.22/72 + (p.space_after.pt/72 if p.space_after else 0)
    return tot
issues=[]
for i,s in enumerate(P.slides,1):
    blocks=[]
    for sh in s.shapes:
        st=str(sh.shape_type)
        l,t=sh.left/EM, sh.top/EM
        w=(sh.width or 0)/EM
        if 'PICTURE' in st: blocks.append(('PIC','', l,t,w,sh.height/EM))
        elif 'TABLE' in st:
            blocks.append(('TAB','', l,t,w,sum(r.height for r in sh.table.rows)/EM))
        elif sh.has_text_frame and sh.text_frame.text.strip():
            h=est_h(sh,w)
            blocks.append(('TXT', sh.text_frame.text[:34].replace('\n',' '), l,t,w,h))
    for a in blocks:
        if a[3]+a[5] > 6.99 and a[0]!='TXT' or (a[0]=='TXT' and a[3]+a[5]>7.0 and a[3]<6.9):
            issues.append(f'slide {i}: {a[0]} "{a[1]}" bottom={a[3]+a[5]:.2f}')
        for b in blocks:
            if a is b or (a[0]=='TXT' and b[0]=='TXT'): continue
            ox = a[2] < b[2]+b[4]-0.05 and b[2] < a[2]+a[4]-0.05
            oy = a[3] < b[3]+b[5]-0.05 and b[3] < a[3]+a[5]-0.05
            if ox and oy and a[0]=='TXT':
                issues.append(f'slide {i}: TXT "{a[1]}" (y {a[3]:.2f}-{a[3]+a[5]:.2f}) '
                              f'overlaps {b[0]} (y {b[3]:.2f}-{b[3]+b[5]:.2f})')
for x in sorted(set(issues)): print(' -',x)
print(len(set(issues)),'issues')
